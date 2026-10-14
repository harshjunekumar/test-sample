#!/usr/bin/env python3
"""Cover letter - Manager, Workforce Planning. Outputs PDF (reportlab) and editable DOCX (python-docx)."""
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY
import docx
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

BASE = "/home/user/test-sample/Juhi_Bhalla_Cover_Letter_Workforce_Planning"
LINKEDIN = "https://www.linkedin.com/in/juhi-bhalla"
DATE = "October 2026"

PARAS = [
  "Dear Hiring Team,",
  "I am applying for the Manager, Workforce Planning role. Over the past five years at Flipkart and Enverus, my work has centred "
  "on the same question your team answers every day: how do we forecast demand accurately, plan capacity to meet it, and keep "
  "performance on track when reality does not follow the plan?",
  "As Category Manager for Large Appliances at Flipkart, I owned forecasting and planning for one of the platform's largest "
  "categories. I built regression-based demand forecasts and turned them into supply and in-stock plans for peak events like Big "
  "Billion Days, where volume spikes leave no room for error. That planning contributed to a <b>55% category revenue uplift</b> during "
  "the event. I owned KPIs across GMV, conversion and in-stock, ran root-cause analysis when numbers slipped, and led recovery "
  "actions that grew category revenue <b>19% year on year</b>. I also held sellers and vendors accountable to their commitments, which "
  "maps closely to managing outsourced partners against staffing and service-level targets.",
  "I care about the mechanics behind a plan. I built and automated Tableau and Power BI dashboards and templates in SQL so leaders "
  "had real-time visibility instead of manual reports, and at Enverus I re-engineered global reporting through Lean Six Sigma, "
  "cutting cycle times by <b>20%</b>. Whether the role leans toward the business-facing track or the technical track, I would bring "
  "analytical rigour and a habit of improving the tools themselves.",
  "I want to be open about where I am growing: my planning experience comes from e-commerce rather than a contact-centre WFM "
  "function, and I am eager to go deep on service level, occupancy, shrinkage and your WFM platforms. The fundamentals of "
  "forecasting, capacity and KPI recovery carry over directly, and I learn new domains fast. I have also led teams on large-scale "
  "events of 800 to 1,000+ participants, handling staffing, scheduling and live issues on the day, and I am excited to grow as a "
  "people leader who coaches a team while staying close to the work.",
  "I would welcome the chance to discuss how I can help your team keep staffing aligned with business priorities. Thank you for "
  "your time and consideration.",
]

# ---------- PDF ----------
BLACK = colors.HexColor("#1a1a1a"); GREY = colors.HexColor("#6b6b6b"); RULE = colors.HexColor("#b7b7b7")
name = ParagraphStyle("n", fontName="Helvetica-Bold", fontSize=16, textColor=BLACK, leading=19)
contact = ParagraphStyle("c", fontName="Helvetica", fontSize=9, textColor=BLACK, leading=12)
meta = ParagraphStyle("m", fontName="Helvetica", fontSize=9.5, textColor=GREY, leading=13)
body = ParagraphStyle("b", fontName="Helvetica", fontSize=10.2, textColor=BLACK, leading=14.6, alignment=TA_JUSTIFY, spaceAfter=9)
doc = SimpleDocTemplate(BASE + ".pdf", pagesize=A4, topMargin=18*mm, bottomMargin=16*mm, leftMargin=20*mm, rightMargin=20*mm,
                        title="Juhi Bhalla - Cover Letter", author="Juhi Bhalla")
E = [Paragraph("JUHI BHALLA", name), Spacer(1, 3),
     Paragraph(f"+91 88709 52224 &nbsp;|&nbsp; juhibhalla.jblko@gmail.com &nbsp;|&nbsp; "
               f"<link href='{LINKEDIN}' color='#0563C1'><u>LinkedIn</u></link> &nbsp;|&nbsp; Bengaluru, India", contact),
     HRFlowable(width="100%", thickness=0.7, color=RULE, spaceBefore=6, spaceAfter=10),
     Paragraph(DATE, meta), Paragraph("Re: Manager, Workforce Planning", meta), Spacer(1, 10)]
for p in PARAS:
    E.append(Paragraph(p, body))
E += [Spacer(1, 4), Paragraph("Sincerely,", body), Paragraph("<b>Juhi Bhalla</b>", body)]
doc.build(E)

# ---------- DOCX ----------
d = docx.Document()
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for s in d.sections:
    s.left_margin = s.right_margin = docx.shared.Mm(20); s.top_margin = docx.shared.Mm(18); s.bottom_margin = docx.shared.Mm(16)
def para(text, bold=False, size=None, color=None, after=8):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    parts = text.replace("<b>", "\x00").replace("</b>", "\x00").split("\x00")
    for i, t in enumerate(parts):
        r = p.add_run(t); r.bold = bold or (i % 2 == 1)
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor.from_string(color)
    return p
para("JUHI BHALLA", bold=True, size=16, after=2)
para("+91 88709 52224  |  juhibhalla.jblko@gmail.com  |  linkedin.com/in/juhi-bhalla  |  Bengaluru, India", size=9.5, after=12)
para(DATE, color="6B6B6B", after=0); para("Re: Manager, Workforce Planning", color="6B6B6B", after=12)
for p in PARAS:
    para(p)
para("Sincerely,", after=0); para("Juhi Bhalla", bold=True)
d.save(BASE + ".docx")
print("built", BASE + ".pdf", BASE + ".docx")
