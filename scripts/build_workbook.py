"""Build the Checkpoint 2/3 Excel workbook for the Flipkart customer-service project.

Every summary number in the workbook is an Excel formula that reads from the
Clean_Data sheet, so the workbook recalculates if the data changes.

Usage: python scripts/build_workbook.py
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "flipkart_customer_calls_raw.csv"
CLEAN_CSV = ROOT / "data" / "flipkart_customer_calls_clean.csv"
OUT = ROOT / "excel" / "Flipkart_Customer_Service_EDA.xlsx"
TREE_PNG = ROOT / "docs" / "metric_tree.png"

BLUE = "2874F0"
NAVY = "172337"
YELLOW = "FFE11B"
LIGHT = "EAF1FE"
FONT = "Arial"

HDR_FILL = PatternFill("solid", fgColor=BLUE)
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT, bold=True, size=16, color=NAVY)
SUB_FONT = Font(name=FONT, bold=True, size=12, color=BLUE)
NOTE_FONT = Font(name=FONT, italic=True, size=9, color="595959")
BOLD = Font(name=FONT, bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
KPI_FILL = PatternFill("solid", fgColor=LIGHT)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SENTIMENTS = ["Very Negative", "Negative", "Neutral", "Positive", "Very Positive"]
RESPONSE = ["Below SLA", "Within SLA", "Above SLA"]
CENTERS = ["Delhi", "Mumbai", "Kolkata", "Chennai"]
CHANNELS = ["Call-Center", "Chatbot", "Email", "Web"]
REASONS = ["Billing Question", "Payments", "Service Outage"]
DUR_BANDS = ["05-15 min", "16-25 min", "26-35 min", "36-45 min"]
CSAT_BANDS = ["Low (1-4)", "Mid (5-7)", "High (8-10)", "No Response"]
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


# ---------------------------------------------------------------- cleaning
def clean(raw: pd.DataFrame) -> tuple[pd.DataFrame, list[tuple]]:
    log = []
    df = raw.copy()
    log.append(("Rows imported", len(df), "Raw CSV loaded without modification into Raw_Data."))

    dup_rows = int(df.duplicated().sum())
    dup_ids = int(df["id"].duplicated().sum())
    df = df.drop_duplicates()
    log.append(("Exact duplicate rows removed", dup_rows, "Checked all 13 columns."))
    log.append(("Duplicate call IDs", dup_ids, "Every call id is unique."))

    for c in df.select_dtypes(["object", "str"]):
        df[c] = df[c].str.strip()

    df = df.rename(columns={"Gender": "gender", "call duration in minutes": "call_duration_min"})

    n = int(df["customer_name"].isna().sum())
    df["customer_name"] = df["customer_name"].fillna("Unknown").str.title()
    log.append(("customer_name missing -> 'Unknown'", n, "Names also converted to Title Case."))

    n = int(df["gender"].isna().sum())
    df["gender"] = df["gender"].map({"f": "Female", "m": "Male"}).fillna("Unknown")
    log.append(("gender missing -> 'Unknown'", n, "Codes f/m standardised to Female/Male."))

    n = int(df["city"].isna().sum())
    df[["city", "state"]] = df[["city", "state"]].fillna("Unknown")
    log.append(("city/state missing -> 'Unknown'", n, "Same rows are missing both fields; kept for analysis."))

    n = int(df["csat_score"].isna().sum())
    log.append((
        "csat_score missing (left blank)", n,
        "Not imputed: a blank means the customer did not answer the survey. "
        "Non-response is itself a metric (CSAT response rate). Averages use rated calls only.",
    ))

    df["call_date"] = pd.to_datetime(df["call_timestamp"], format="%m/%d/%Y")
    df = df.drop(columns="call_timestamp")
    log.append(("call_timestamp (text MM/DD/YYYY) -> call_date (real date)", len(df),
                f"Range {df.call_date.min():%d-%b-%Y} to {df.call_date.max():%d-%b-%Y}."))

    bad = int((~df["call_duration_min"].between(1, 180)).sum())
    log.append(("call_duration_min outside 1-180 min", bad, "No outliers; range is 5-45 minutes."))
    log.append(("Rows in Clean_Data", len(df), ""))

    cols = ["id", "customer_name", "gender", "sentiment", "csat_score", "call_date", "reason",
            "city", "state", "channel", "response_time", "call_duration_min", "call_center"]
    return df[cols].reset_index(drop=True), log


# ---------------------------------------------------------------- helpers
def style_header(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row, c)
        cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def box(ws, r1, c1, r2, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).border = BOX


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    if sub:
        ws["A2"] = sub
        ws["A2"].font = NOTE_FONT
    ws.sheet_view.showGridLines = False


def main():
    raw = pd.read_csv(RAW)
    df, log = clean(raw)
    n = len(df)
    last = n + 1  # last data row in Clean_Data

    wb = Workbook()
    wb._named_styles["Normal"].font = Font(name=FONT, size=10)

    # Clean_Data column letters
    C = {k: v for k, v in zip(
        ["id", "customer_name", "gender", "sentiment", "csat", "date", "reason", "city", "state",
         "channel", "resp", "dur", "center", "day_num", "day", "sent_score", "responded",
         "csat_band", "sla_breach", "resp_rank", "dur_band", "csat_sq", "satisfied", "neg"],
        [get_column_letter(i) for i in range(1, 25)])}

    def rng(key):
        return f"Clean_Data!${C[key]}$2:${C[key]}${last}"

    # ------------------------------------------------------------ README
    ws = wb.active
    ws.title = "README"
    title(ws, "Flipkart Customer Service & Retention: EDA Workbook",
          "Data: 30,000 customer-service contacts, 1-31 Oct 2020. Every summary is a live formula on Clean_Data.")
    rows = [
        ("Sheet", "What it contains", "Checkpoint"),
        ("Metrics_Hypotheses", "Key metrics, metric tree, hypotheses and the columns used to test each one", "1"),
        ("Raw_Data", "The CSV exactly as provided (untouched)", "2"),
        ("Cleaning_Log", "Every cleaning step with the number of rows affected", "2"),
        ("Clean_Data", "Cleaned data + 11 derived columns (dark headers N-X)", "2"),
        ("Descriptive_Stats", "Mean, median, std dev, quartiles, KPIs", "2"),
        ("Correlation", "Correlation matrix + CSAT spread by location / channel", "2"),
        ("Pivots", "Pivot-style summaries by sentiment, SLA, call centre, channel, reason, day, date", "2"),
        ("Charts", "Visualisations built on the Pivots sheet", "2"),
        ("Hypothesis_Tests", "Statistical test and verdict for every hypothesis", "2"),
        ("Dashboard", "Interactive dashboard: pick call centre / channel / reason in the yellow cells", "3"),
    ]
    for i, r in enumerate(rows, start=4):
        for j, v in enumerate(r, start=1):
            ws.cell(i, j, v)
    style_header(ws, 4, 1, 3)
    box(ws, 5, 1, 4 + len(rows) - 1, 3)
    ws["A17"] = "Colour legend"
    ws["A17"].font = SUB_FONT
    ws["A18"] = "Yellow cell"
    ws["A18"].fill = INPUT_FILL
    ws["B18"] = "Input you can change (Dashboard filters)"
    ws["A19"] = "Light blue cell"
    ws["A19"].fill = KPI_FILL
    ws["B19"] = "Key metric / result"
    ws["A21"] = ("Pivot note: the Pivots sheet uses COUNTIFS / AVERAGEIFS so it recalculates in Excel, "
                 "LibreOffice and Google Sheets. To add a native PivotTable, select Clean_Data and use "
                 "Insert > PivotTable. The fields to use are the same.")
    ws["A21"].font = NOTE_FONT
    widths(ws, {"A": 22, "B": 90, "C": 12})

    # ------------------------------------------------------------ Metrics_Hypotheses
    ws = wb.create_sheet("Metrics_Hypotheses")
    title(ws, "Checkpoint 1: Metrics, Metric Tree and Hypotheses")
    ws["A3"] = "North-star goal: improve customer retention. Retention itself is not in the data, so CSAT is the proxy outcome."
    ws["A3"].font = BOLD
    metrics = [
        ("Level", "Metric", "Definition / formula", "Why it matters for retention", "Column(s)"),
        ("Outcome (L1)", "Average CSAT", "Mean csat_score of rated calls (1-10)", "Main proxy for satisfaction and retention", "csat_score"),
        ("Outcome (L1)", "% Satisfied (CSAT >= 8)", "Rated calls with CSAT 8-10 / rated calls", "Satisfied customers are the most likely to come back", "csat_score"),
        ("Outcome (L1)", "% Dissatisfied (CSAT <= 4)", "Rated calls with CSAT 1-4 / rated calls", "Early churn-risk signal", "csat_score"),
        ("Outcome (L1)", "CSAT response rate", "Rated calls / all calls", "Low response = weak feedback loop and a hidden risk", "csat_score"),
        ("Experience (L2)", "Negative sentiment share", "(Negative + Very Negative) / calls", "Emotional state of the customer at contact", "sentiment"),
        ("Experience (L2)", "Avg sentiment score", "Very Negative=1 ... Very Positive=5", "Numeric version of sentiment for correlation", "sentiment"),
        ("Speed (L3)", "SLA compliance", "(Within + Below SLA) / calls", "Customers dislike waiting", "response_time"),
        ("Speed (L3)", "SLA breach rate", "Above SLA / calls", "Direct measure of slow responses", "response_time"),
        ("Efficiency (L3)", "Avg call duration (AHT)", "Mean call_duration_min", "Long calls can mean hard-to-solve problems", "call duration in minutes"),
        ("Demand (L3)", "Contact volume & mix", "Calls by reason / channel / date", "High contact volume shows friction in the product", "reason, channel, call_timestamp"),
        ("Operations (L3)", "Location performance", "All metrics above split by call_center", "Finds weak sites or training gaps", "call_center, city, state"),
    ]
    for i, r in enumerate(metrics, start=5):
        for j, v in enumerate(r, start=1):
            ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
    style_header(ws, 5, 1, 5)
    box(ws, 6, 1, 4 + len(metrics), 5)

    r0 = 5 + len(metrics) + 2
    ws.cell(r0, 1, "Hypotheses").font = SUB_FONT
    hyps = [
        ("ID", "Hypothesis", "Metric tested", "Columns", "Test used"),
        ("H1", "Reducing response time (fewer Above-SLA contacts) will improve CSAT and retention",
         "Avg CSAT by response_time", "response_time, csat_score", "Welch t-test: Above SLA vs rest"),
        ("H2", "Customer-service training to address negative sentiment will increase CSAT and retention",
         "Avg CSAT by sentiment", "sentiment, csat_score", "Pearson r (sentiment score vs CSAT)"),
        ("H3", "Reducing call duration through efficient processes will improve CSAT and retention",
         "Avg CSAT by duration band", "call duration in minutes, csat_score", "Pearson r + t-test on r"),
        ("H4", "CSAT and SLA performance differ by call-centre location",
         "Avg CSAT / SLA breach by call_center", "call_center, csat_score, response_time", "Range between centres"),
        ("H5", "Some channels (e.g. Chatbot) give lower CSAT than others",
         "Avg CSAT by channel", "channel, csat_score", "Welch t-test: Chatbot vs rest"),
        ("H6", "Unhappy customers skip the CSAT survey (non-response bias)",
         "Response rate by sentiment", "sentiment, csat_score", "Response-rate gap Neg vs Pos"),
    ]
    for i, r in enumerate(hyps, start=r0 + 1):
        for j, v in enumerate(r, start=1):
            ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
    style_header(ws, r0 + 1, 1, 5)
    box(ws, r0 + 2, 1, r0 + len(hyps), 5)
    widths(ws, {"A": 16, "B": 48, "C": 36, "D": 38, "E": 32})

    tr = r0 + len(hyps) + 2
    ws.cell(tr, 1, "Metric tree").font = SUB_FONT
    if TREE_PNG.exists():
        img = Image(str(TREE_PNG))
        img.width, img.height = 1100, int(1100 * img.height / img.width)
        ws.add_image(img, f"A{tr + 1}")

    # ------------------------------------------------------------ Raw_Data
    ws = wb.create_sheet("Raw_Data")
    ws.append(list(raw.columns))
    for rec in raw.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in rec])
    style_header(ws, 1, 1, raw.shape[1])
    ws.freeze_panes = "A2"
    for i in range(1, raw.shape[1] + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16

    # ------------------------------------------------------------ Cleaning_Log
    ws = wb.create_sheet("Cleaning_Log")
    title(ws, "Data cleaning & preparation log", "Raw_Data is kept untouched; all steps below produce Clean_Data.")
    ws.append([])
    ws.append(["Step", "Rows affected", "Notes"])
    style_header(ws, 4, 1, 3)
    for step in log:
        ws.append(list(step))
    box(ws, 5, 1, 4 + len(log), 3)
    for r in range(5, 5 + len(log)):
        ws.cell(r, 2).number_format = "#,##0"
        ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
    r = 6 + len(log)
    ws.cell(r, 1, "Derived columns added to Clean_Data (Excel formula equivalent shown; values computed in data prep)").font = SUB_FONT
    derived = [
        ("day_num", "=WEEKDAY(call_date,2)", "1 = Monday ... 7 = Sunday"),
        ("day_name", '=TEXT(call_date,"ddd")', "Mon ... Sun"),
        ("sentiment_score", "=MATCH(sentiment,{Very Negative..Very Positive},0)", "1-5 ordinal scale"),
        ("csat_responded", '=IF(csat_score="",0,1)', "1 if customer rated the call"),
        ("csat_band", "Low 1-4 / Mid 5-7 / High 8-10 / No Response", "Satisfaction group"),
        ("sla_breach", '=IF(response_time="Above SLA",1,0)', "SLA breach flag"),
        ("response_rank", "Below=1, Within=2, Above=3", "Ordinal response speed"),
        ("duration_band", "5-15 / 16-25 / 26-35 / 36-45 min", "Call-length group"),
        ("csat_sq", "=csat_score^2", "Helper for variance in t-tests"),
        ("satisfied", "=1 if CSAT >= 8 (blank if not rated)", "Satisfied flag"),
        ("negative_sentiment", "=1 if sentiment_score <= 2", "Negative sentiment flag"),
    ]
    ws.cell(r + 1, 1, "Column")
    ws.cell(r + 1, 2, "Logic")
    ws.cell(r + 1, 3, "Meaning")
    style_header(ws, r + 1, 1, 3)
    for i, d in enumerate(derived, start=r + 2):
        for j, v in enumerate(d, start=1):
            ws.cell(i, j, v).data_type = "s"  # show formula text literally
    box(ws, r + 2, 1, r + 1 + len(derived), 3)
    widths(ws, {"A": 52, "B": 52, "C": 80})

    # ------------------------------------------------------------ Clean_Data
    ws = wb.create_sheet("Clean_Data")
    headers = list(df.columns) + ["day_num", "day_name", "sentiment_score", "csat_responded", "csat_band",
                                  "sla_breach", "response_rank", "duration_band", "csat_sq", "satisfied",
                                  "negative_sentiment"]
    ws.append(headers)
    # Derived columns are computed here (Excel equivalents are listed on Cleaning_Log).
    # Writing 330k row-level formulas makes the workbook too slow to recalculate.
    sent_score = df["sentiment"].map({s: i for i, s in enumerate(SENTIMENTS, start=1)})
    csat = df["csat_score"]
    dur = df["call_duration_min"]
    derived_df = pd.DataFrame({
        "day_num": df["call_date"].dt.dayofweek + 1,
        "day_name": df["call_date"].dt.strftime("%a"),
        "sentiment_score": sent_score,
        "csat_responded": csat.notna().astype(int),
        "csat_band": pd.cut(csat, [0, 4, 7, 10], labels=CSAT_BANDS[:3]).astype(object).where(csat.notna(), "No Response"),
        "sla_breach": (df["response_time"] == "Above SLA").astype(int),
        "response_rank": df["response_time"].map({s: i for i, s in enumerate(RESPONSE, start=1)}),
        "duration_band": pd.cut(dur, [0, 15, 25, 35, 45], labels=DUR_BANDS).astype(str),
        "csat_sq": csat ** 2,
        "satisfied": (csat >= 8).astype(float).where(csat.notna()),
        "negative_sentiment": (sent_score <= 2).astype(int),
    })
    full = pd.concat([df, derived_df], axis=1)
    full.to_csv(CLEAN_CSV, index=False)
    for rec in full.itertuples(index=False):
        vals = [None if (isinstance(v, float) and pd.isna(v)) else v for v in rec]
        vals[5] = vals[5].to_pydatetime()
        ws.append(vals)
    style_header(ws, 1, 1, len(headers))
    for c in range(14, 25):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=NAVY)
    for r in range(2, last + 1):
        ws.cell(r, 6).number_format = "dd-mmm-yyyy"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:X{last}"
    for i in range(1, 25):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.column_dimensions["A"].width = 27

    # ------------------------------------------------------------ Descriptive_Stats
    ws = wb.create_sheet("Descriptive_Stats")
    title(ws, "Descriptive statistics", "CSAT statistics use rated calls only (blank = no survey answer).")
    ws.append([])
    ws.append(["Statistic", "CSAT score", "Call duration (min)", "Sentiment score (1-5)", "Response rank (1-3)"])
    style_header(ws, 4, 1, 5)
    stat_rows = [
        ("Count (non-blank)", "COUNT"), ("Mean", "AVERAGE"), ("Median", "MEDIAN"), ("Mode", "MODE"),
        ("Standard deviation", "STDEV"), ("Variance", "VAR"), ("Minimum", "MIN"),
        ("25th percentile", "Q1"), ("75th percentile", "Q3"), ("Maximum", "MAX"),
        ("Skewness", "SKEW"), ("Coefficient of variation", "CV"),
    ]
    cols = [rng("csat"), rng("dur"), rng("sent_score"), rng("resp_rank")]
    for i, (label, fn) in enumerate(stat_rows, start=5):
        ws.cell(i, 1, label)
        for j, rg in enumerate(cols, start=2):
            if fn == "Q1":
                f = f"=QUARTILE({rg},1)"
            elif fn == "Q3":
                f = f"=QUARTILE({rg},3)"
            elif fn == "CV":
                f = f"={get_column_letter(j)}9/{get_column_letter(j)}6"
            else:
                f = f"={fn}({rg})"
            ws.cell(i, j, f).number_format = "#,##0" if fn == "COUNT" else ("0.0%" if fn == "CV" else "0.00")
    box(ws, 5, 1, 16, 5)

    ws["A19"] = "Headline KPIs"
    ws["A19"].font = SUB_FONT
    ws.append(["KPI", "Value", "Formula logic"])
    style_header(ws, 20, 1, 3)
    kpis = [
        ("Total contacts", f"=COUNTA({rng('id')})", "#,##0", "Rows in Clean_Data"),
        ("Rated contacts", f"=SUM({rng('responded')})", "#,##0", "Calls with a CSAT score"),
        ("CSAT response rate", "=B22/B21", "0.0%", "Rated / total"),
        ("Average CSAT (1-10)", f"=AVERAGE({rng('csat')})", "0.00", "Rated calls only"),
        ("% Satisfied (CSAT >= 8)", f'=COUNTIF({rng("csat")},">=8")/B22', "0.0%", "Share of rated calls"),
        ("% Dissatisfied (CSAT <= 4)", f'=COUNTIF({rng("csat")},"<=4")/B22', "0.0%", "Share of rated calls"),
        ("Net satisfaction (Sat - Dissat)", "=B25-B26", "0.0%", "Positive is good"),
        ("Negative sentiment share", f"=SUM({rng('neg')})/B21", "0.0%", "Negative + Very Negative"),
        ("SLA breach rate (Above SLA)", f"=SUM({rng('sla_breach')})/B21", "0.0%", "Slow responses"),
        ("SLA compliance", "=1-B29", "0.0%", "Within + Below SLA"),
        ("Avg call duration (min)", f"=AVERAGE({rng('dur')})", "0.0", "All calls"),
        ("Billing share of contacts", f'=COUNTIF({rng("reason")},"Billing Question")/B21', "0.0%", "Largest contact driver"),
    ]
    for i, (k, f, fmt, note) in enumerate(kpis, start=21):
        ws.cell(i, 1, k)
        c = ws.cell(i, 2, f)
        c.number_format, c.fill, c.font = fmt, KPI_FILL, BOLD
        ws.cell(i, 3, note).font = NOTE_FONT
    box(ws, 21, 1, 20 + len(kpis), 3)

    ws["A35"] = "CSAT frequency distribution"
    ws["A35"].font = SUB_FONT
    ws.append(["CSAT score", "Calls", "% of rated"])
    style_header(ws, 36, 1, 3)
    for s in range(1, 11):
        r = 36 + s
        ws.cell(r, 1, s)
        ws.cell(r, 2, f"=COUNTIF({rng('csat')},A{r})").number_format = "#,##0"
        ws.cell(r, 3, f"=B{r}/$B$22").number_format = "0.0%"
    box(ws, 37, 1, 46, 3)
    widths(ws, {"A": 32, "B": 18, "C": 22, "D": 22, "E": 22})

    ch = BarChart()
    ch.title, ch.y_axis.title, ch.x_axis.title = "CSAT score distribution (rated calls)", "Calls", "CSAT"
    ch.add_data(Reference(ws, min_col=2, min_row=36, max_row=46), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=37, max_row=46))
    ch.legend, ch.height, ch.width = None, 7.5, 16
    ws.add_chart(ch, "E35")

    # ------------------------------------------------------------ Correlation
    ws = wb.create_sheet("Correlation")
    title(ws, "Correlation analysis",
          "Pearson r, pairwise on rated calls. |r| < 0.1 = no meaningful relationship; > 0.7 = strong.")
    vars_ = [("CSAT score", "csat"), ("Sentiment score", "sent_score"), ("Call duration", "dur"),
             ("Response rank", "resp_rank"), ("SLA breach flag", "sla_breach"), ("Day of week", "day_num")]
    ws.cell(4, 1, "Variable")
    for j, (lab, _) in enumerate(vars_, start=2):
        ws.cell(4, j, lab)
    style_header(ws, 4, 1, len(vars_) + 1)
    for i, (lab_i, k_i) in enumerate(vars_, start=5):
        ws.cell(i, 1, lab_i).font = BOLD
        for j, (_, k_j) in enumerate(vars_, start=2):
            ws.cell(i, j, f"=CORREL({rng(k_i)},{rng(k_j)})").number_format = "0.000"
    box(ws, 5, 1, 4 + len(vars_), len(vars_) + 1)
    ws.conditional_formatting.add(
        f"B5:{get_column_letter(len(vars_) + 1)}{4 + len(vars_)}",
        ColorScaleRule(start_type="num", start_value=-1, start_color="F8696B",
                       mid_type="num", mid_value=0, mid_color="FFFFFF",
                       end_type="num", end_value=1, end_color="63BE7B"))

    ws["A13"] = "Significance of each driver's correlation with CSAT"
    ws["A13"].font = SUB_FONT
    ws.append(["Driver", "r with CSAT", "n (rated)", "t statistic", "p-value", "R-squared", "Reading"])
    style_header(ws, 14, 1, 7)
    for i, (lab, _) in enumerate(vars_[1:], start=15):
        src = 6 + (i - 15)
        ws.cell(i, 1, lab)
        ws.cell(i, 2, f"=B{src}").number_format = "0.000"
        ws.cell(i, 3, "=Descriptive_Stats!$B$22").number_format = "#,##0"
        ws.cell(i, 4, f"=B{i}*SQRT((C{i}-2)/(1-B{i}^2))").number_format = "0.00"
        ws.cell(i, 5, f"=TDIST(ABS(D{i}),C{i}-2,2)").number_format = "0.0000"
        ws.cell(i, 6, f"=B{i}^2").number_format = "0.0%"
        ws.cell(i, 7, f'=IF(E{i}>=0.05,"No significant link",IF(ABS(B{i})>=0.7,"Strong driver",'
                      f'IF(ABS(B{i})>=0.3,"Moderate driver","Significant but negligible")))')
        ws.cell(i, 7).fill = KPI_FILL
    box(ws, 15, 1, 19, 7)

    ws["A22"] = "CSAT spread by location and channel (categorical relationships)"
    ws["A22"].font = SUB_FONT
    ws.append(["Dimension", "Highest avg CSAT", "Lowest avg CSAT", "Spread (points)", "Overall CSAT", "Reading"])
    style_header(ws, 23, 1, 6)
    spreads = [("Call centre", "Pivots!$C$15:$C$18"), ("Channel", "Pivots!$C$24:$C$27"),
               ("Contact reason", "Pivots!$C$33:$C$35"), ("Response time (SLA)", "Pivots!$C$41:$C$43"),
               ("Sentiment", "Pivots!$C$5:$C$9")]
    for i, (lab, rg) in enumerate(spreads, start=24):
        ws.cell(i, 1, lab)
        ws.cell(i, 2, f"=MAX({rg})").number_format = "0.00"
        ws.cell(i, 3, f"=MIN({rg})").number_format = "0.00"
        ws.cell(i, 4, f"=B{i}-C{i}").number_format = "0.00"
        ws.cell(i, 5, "=Descriptive_Stats!$B$24").number_format = "0.00"
        ws.cell(i, 6, f'=IF(D{i}<0.3,"Negligible difference",IF(D{i}<1,"Small difference","Large difference"))')
        ws.cell(i, 6).fill = KPI_FILL
    box(ws, 24, 1, 28, 6)
    widths(ws, {"A": 26, "B": 18, "C": 18, "D": 18, "E": 16, "F": 22, "G": 26})

    # ------------------------------------------------------------ Pivots
    ws = wb.create_sheet("Pivots")
    title(ws, "Pivot summaries", "Formula-based pivots (COUNTIFS / AVERAGEIFS) on Clean_Data.")

    def pivot(r, label, key, items):
        ws.cell(r - 1, 1, f"By {label}").font = SUB_FONT
        hdr = [label, "Calls", "Avg CSAT", "% of calls", "Response rate", "% Satisfied (>=8)",
               "SLA breach %", "Avg duration", "Negative sentiment %"]
        for j, h in enumerate(hdr, start=1):
            ws.cell(r, j, h)
        style_header(ws, r, 1, len(hdr))
        for i, it in enumerate(items, start=r + 1):
            k = rng(key)
            ws.cell(i, 1, it).font = BOLD
            ws.cell(i, 2, f"=COUNTIFS({k},A{i})").number_format = "#,##0"
            ws.cell(i, 3, f'=IFERROR(AVERAGEIFS({rng("csat")},{k},A{i}),"")').number_format = "0.00"
            ws.cell(i, 4, f"=B{i}/Descriptive_Stats!$B$21").number_format = "0.0%"
            ws.cell(i, 5, f"=SUMIFS({rng('responded')},{k},A{i})/B{i}").number_format = "0.0%"
            ws.cell(i, 6, f'=IFERROR(COUNTIFS({k},A{i},{rng("csat")},">=8")/SUMIFS({rng("responded")},{k},A{i}),"")').number_format = "0.0%"
            ws.cell(i, 7, f"=SUMIFS({rng('sla_breach')},{k},A{i})/B{i}").number_format = "0.0%"
            ws.cell(i, 8, f"=AVERAGEIFS({rng('dur')},{k},A{i})").number_format = "0.0"
            ws.cell(i, 9, f"=SUMIFS({rng('neg')},{k},A{i})/B{i}").number_format = "0.0%"
        box(ws, r + 1, 1, r + len(items), len(hdr))
        return r + 1, r + len(items)

    # fixed positions (referenced by Correlation / Charts / Hypothesis_Tests)
    P = {
        "sent": pivot(4, "sentiment", "sentiment", SENTIMENTS),        # rows 5-9
        "center": pivot(14, "call_center", "center", CENTERS),         # rows 15-18
        "channel": pivot(23, "channel", "channel", CHANNELS),          # rows 24-27
        "reason": pivot(32, "reason", "reason", REASONS),              # rows 33-35
        "resp": pivot(40, "response_time", "resp", RESPONSE),          # rows 41-43
        "dur": pivot(47, "duration_band", "dur_band", DUR_BANDS),      # rows 48-51
        "day": pivot(55, "day_name", "day", DAYS),                     # rows 56-62
    }

    ws.cell(63, 1, "Oct 2020 has 5 Thursdays, Fridays and Saturdays (31 Oct holds 1 contact), so raw day totals differ. Per day, volume is flat at ~1,000 contacts.").font = NOTE_FONT

    # Call centre x response time (row %)
    r = 66
    ws.cell(r - 1, 1, "Call centre x response time (% of the centre's calls)").font = SUB_FONT
    ws.cell(r, 1, "call_center")
    for j, s in enumerate(RESPONSE, start=2):
        ws.cell(r, j, s)
    style_header(ws, r, 1, 4)
    for i, c in enumerate(CENTERS, start=r + 1):
        ws.cell(i, 1, c).font = BOLD
        for j in range(2, 5):
            col = get_column_letter(j)
            ws.cell(i, j, f"=COUNTIFS({rng('center')},$A{i},{rng('resp')},{col}${r})/COUNTIFS({rng('center')},$A{i})").number_format = "0.0%"
    box(ws, r + 1, 1, r + 4, 4)

    # Sentiment x CSAT band (counts)
    r = 74
    ws.cell(r - 1, 1, "Sentiment x CSAT band (calls)").font = SUB_FONT
    ws.cell(r, 1, "sentiment")
    for j, s in enumerate(CSAT_BANDS, start=2):
        ws.cell(r, j, s)
    style_header(ws, r, 1, 5)
    for i, s in enumerate(SENTIMENTS, start=r + 1):
        ws.cell(i, 1, s).font = BOLD
        for j in range(2, 6):
            col = get_column_letter(j)
            ws.cell(i, j, f"=COUNTIFS({rng('sentiment')},$A{i},{rng('csat_band')},{col}${r})").number_format = "#,##0"
    box(ws, r + 1, 1, r + 5, 5)
    ws.conditional_formatting.add(f"B{r+1}:D{r+5}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color=BLUE))

    # Call centre x channel avg CSAT
    r = 83
    ws.cell(r - 1, 1, "Avg CSAT: call centre x channel").font = SUB_FONT
    ws.cell(r, 1, "call_center")
    for j, s in enumerate(CHANNELS, start=2):
        ws.cell(r, j, s)
    style_header(ws, r, 1, 5)
    for i, c in enumerate(CENTERS, start=r + 1):
        ws.cell(i, 1, c).font = BOLD
        for j in range(2, 6):
            col = get_column_letter(j)
            ws.cell(i, j, f"=AVERAGEIFS({rng('csat')},{rng('center')},$A{i},{rng('channel')},{col}${r})").number_format = "0.00"
    box(ws, r + 1, 1, r + 4, 5)
    ws.conditional_formatting.add(f"B{r+1}:E{r+4}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))

    # Daily trend
    r = 91
    ws.cell(r - 1, 1, "Daily trend").font = SUB_FONT
    for j, h in enumerate(["call_date", "Calls", "Avg CSAT", "SLA breach %", "Negative sentiment %"], start=1):
        ws.cell(r, j, h)
    style_header(ws, r, 1, 5)
    dates = sorted(df["call_date"].dt.date.unique())
    for i, d in enumerate(dates, start=r + 1):
        ws.cell(i, 1, d).number_format = "dd-mmm"
        ws.cell(i, 2, f"=COUNTIFS({rng('date')},A{i})").number_format = "#,##0"
        ws.cell(i, 3, f'=IFERROR(AVERAGEIFS({rng("csat")},{rng("date")},A{i}),"")').number_format = "0.00"
        ws.cell(i, 4, f"=SUMIFS({rng('sla_breach')},{rng('date')},A{i})/B{i}").number_format = "0.0%"
        ws.cell(i, 5, f"=SUMIFS({rng('neg')},{rng('date')},A{i})/B{i}").number_format = "0.0%"
    daily_end = r + len(dates)
    ws.cell(daily_end + 1, 1, "31-Oct has only 1 contact (no CSAT rating), so it is a partial day and is left out of trend reading.").font = NOTE_FONT
    box(ws, r + 1, 1, daily_end, 5)

    # Top states
    r = daily_end + 3
    ws.cell(r - 1, 1, "Top 10 states by contact volume").font = SUB_FONT
    for j, h in enumerate(["state", "Calls", "Avg CSAT", "SLA breach %"], start=1):
        ws.cell(r, j, h)
    style_header(ws, r, 1, 4)
    top_states = df["state"].value_counts().head(10).index
    for i, s in enumerate(top_states, start=r + 1):
        ws.cell(i, 1, s).font = BOLD
        ws.cell(i, 2, f"=COUNTIFS({rng('state')},A{i})").number_format = "#,##0"
        ws.cell(i, 3, f"=AVERAGEIFS({rng('csat')},{rng('state')},A{i})").number_format = "0.00"
        ws.cell(i, 4, f"=SUMIFS({rng('sla_breach')},{rng('state')},A{i})/B{i}").number_format = "0.0%"
    box(ws, r + 1, 1, r + 10, 4)
    ws.cell(r + 11, 1, "State list chosen as the 10 largest by volume when the workbook was built.").font = NOTE_FONT
    widths(ws, {get_column_letter(i): 16 for i in range(1, 10)})
    ws.column_dimensions["A"].width = 20
    pv = ws

    # ------------------------------------------------------------ Charts
    ws = wb.create_sheet("Charts")
    title(ws, "Visualisations", "All charts read from the Pivots sheet.")

    def bar(title_, cats, data_cols, r1, r2, anchor, ytitle="Avg CSAT", stacked=False, horiz=False, pct=False):
        ch = BarChart()
        ch.type = "bar" if horiz else "col"
        ch.title, ch.y_axis.title = title_, ytitle
        for dc in data_cols:
            ch.add_data(Reference(pv, min_col=dc, min_row=r1 - 1, max_row=r2), titles_from_data=True)
        ch.set_categories(Reference(pv, min_col=1, min_row=r1, max_row=r2))
        if stacked:
            ch.grouping, ch.overlap = "percentStacked", 100
        if len(data_cols) == 1:
            ch.legend = None
            ch.dataLabels = DataLabelList()
            ch.dataLabels.showVal = True
        if pct:
            ch.y_axis.numFmt = "0%"
        ch.y_axis.delete = False
        ch.x_axis.delete = False
        ch.height, ch.width = 7.5, 15
        ws.add_chart(ch, anchor)

    bar("Avg CSAT by sentiment (H2)", None, [3], 5, 9, "A4")
    bar("Avg CSAT by response time (H1)", None, [3], 41, 43, "J4")
    bar("Avg CSAT by call-duration band (H3)", None, [3], 48, 51, "A20")
    bar("Avg CSAT by call centre (H4)", None, [3], 15, 18, "J20")
    bar("Avg CSAT by channel (H5)", None, [3], 24, 27, "A36")
    bar("CSAT response rate by sentiment (H6)", None, [5], 5, 9, "J36", ytitle="Response rate", pct=True)
    bar("Contact volume by reason", None, [2], 33, 35, "A52", ytitle="Calls")
    bar("Response-time mix by call centre", None, [2, 3, 4], 67, 70, "J52", ytitle="% of calls", stacked=True, horiz=True, pct=True)

    lc = LineChart()
    lc.title, lc.y_axis.title = "Daily contacts and avg CSAT (Oct 2020)", "Calls"
    lc.add_data(Reference(pv, min_col=2, min_row=91, max_row=daily_end), titles_from_data=True)
    lc.set_categories(Reference(pv, min_col=1, min_row=92, max_row=daily_end))
    lc2 = LineChart()
    lc2.add_data(Reference(pv, min_col=3, min_row=91, max_row=daily_end), titles_from_data=True)
    lc2.y_axis.axId, lc2.y_axis.title, lc2.y_axis.crosses = 200, "Avg CSAT", "max"
    lc2.y_axis.scaling.min, lc2.y_axis.scaling.max = 4, 7
    lc.x_axis.number_format = "dd-mmm"
    lc.y_axis.delete = lc.x_axis.delete = lc2.y_axis.delete = False
    lc += lc2
    lc.height, lc.width = 8, 33
    ws.add_chart(lc, "A68")

    pc = PieChart()
    pc.title = "Contact share by channel"
    pc.add_data(Reference(pv, min_col=2, min_row=23, max_row=27), titles_from_data=True)
    pc.set_categories(Reference(pv, min_col=1, min_row=24, max_row=27))
    pc.dataLabels = DataLabelList()
    pc.dataLabels.showPercent = True
    pc.height, pc.width = 7.5, 15
    ws.add_chart(pc, "A85")
    bar("Contacts by day of week", None, [2], 56, 62, "J85", ytitle="Calls")

    # ------------------------------------------------------------ Hypothesis_Tests
    ws = wb.create_sheet("Hypothesis_Tests")
    title(ws, "Hypothesis testing", "alpha = 0.05 (cell B3). Welch t-tests use rated calls only.")
    ws["A3"] = "Significance level (alpha)"
    ws["B3"] = 0.05
    ws["B3"].fill = INPUT_FILL
    ws["B3"].font = Font(name=FONT, color="0000FF")
    ws["B3"].comment = Comment("Standard 5% significance level. Change it to re-run all verdicts.", "Analyst")

    def welch(r, label, key, group_a, a_label, b_label):
        """Welch two-sample t-test: group_a vs everything else on CSAT."""
        k, cs, sq, rs = rng(key), rng("csat"), rng("csat_sq"), rng("responded")
        ws.cell(r, 1, label).font = SUB_FONT
        hdr = ["Group", "n (rated)", "Mean CSAT", "Variance"]
        for j, h in enumerate(hdr, start=1):
            ws.cell(r + 1, j, h)
        style_header(ws, r + 1, 1, 4)
        for i, (lab, crit) in enumerate([(a_label, f'"{group_a}"'), (b_label, f'"<>{group_a}"')], start=r + 2):
            ws.cell(i, 1, lab)
            ws.cell(i, 2, f"=SUMIFS({rs},{k},{crit})").number_format = "#,##0"
            ws.cell(i, 3, f"=AVERAGEIFS({cs},{k},{crit})").number_format = "0.000"
            ws.cell(i, 4, f"=(SUMIFS({sq},{k},{crit})-B{i}*C{i}^2)/(B{i}-1)").number_format = "0.000"
        a, b = r + 2, r + 3
        out = [
            ("Difference in mean CSAT", f"=C{a}-C{b}", "0.000"),
            ("t statistic", f"=(C{a}-C{b})/SQRT(D{a}/B{a}+D{b}/B{b})", "0.00"),
            ("Welch degrees of freedom", f"=(D{a}/B{a}+D{b}/B{b})^2/((D{a}/B{a})^2/(B{a}-1)+(D{b}/B{b})^2/(B{b}-1))", "0"),
            ("p-value (two-tailed)", None, "0.0000"),
        ]
        for i, (lab, f, fmt) in enumerate(out, start=r + 4):
            ws.cell(i, 1, lab)
            if f is None:
                f = f"=TDIST(ABS(B{i-2}),B{i-1},2)"
            ws.cell(i, 2, f).number_format = fmt
        box(ws, r + 2, 1, r + 7, 4)
        return r + 5, r + 7  # t row, p row

    h1_t, h1_p = welch(6, "H1: Response time (Above SLA vs Within/Below SLA)", "resp", "Above SLA", "Above SLA", "Within / Below SLA")
    h5_t, h5_p = welch(16, "H5: Channel (Chatbot vs other channels)", "channel", "Chatbot", "Chatbot", "Other channels")

    ws["A27"] = "Hypothesis verdicts"
    ws["A27"].font = SUB_FONT
    for j, h in enumerate(["ID", "Hypothesis", "Key evidence", "Statistic", "p-value", "Verdict", "Business reading"], start=1):
        ws.cell(28, j, h)
    style_header(ws, 28, 1, 7)
    verdicts = [
        ("H1", "Faster response (fewer Above-SLA) raises CSAT",
         '="Above SLA "&TEXT(C8,"0.00")&" vs rest "&TEXT(C9,"0.00")', f"=B{h1_t}", f"=B{h1_p}",
         '=IF(E29<$B$3,IF(D29<0,"Supported","Rejected (opposite direction)"),"Not supported")',
         "SLA status alone does not move CSAT. Keep SLA as a hygiene metric but it is not the lever."),
        ("H2", "Negative sentiment drives low CSAT",
         '="r = "&TEXT(Correlation!B15,"0.000")&"; Very Neg "&TEXT(Pivots!C5,"0.0")&" vs Very Pos "&TEXT(Pivots!C9,"0.0")',
         "=Correlation!D15", "=Correlation!E15",
         '=IF(AND(E30<$B$3,D30>0),"Supported","Not supported")',
         "Sentiment explains most of the CSAT variance. Improving how agents handle upset customers is the main lever."),
        ("H3", "Shorter calls raise CSAT",
         '="r = "&TEXT(Correlation!B16,"0.000")&"; band spread "&TEXT(MAX(Pivots!C48:C51)-MIN(Pivots!C48:C51),"0.00")',
         "=Correlation!D16", "=Correlation!E16",
         '=IF(AND(E31<$B$3,D31<0),"Supported","Not supported")',
         "Call length is not linked to satisfaction. Do not cut AHT at the cost of resolution."),
        ("H4", "CSAT / SLA differ by call centre",
         '="CSAT spread "&TEXT(Correlation!D24,"0.00")&" pts; SLA breach "&TEXT(MIN(Pivots!G15:G18),"0.0%")&"-"&TEXT(MAX(Pivots!G15:G18),"0.0%")',
         "=Correlation!D24", "n/a",
         '=IF(Correlation!D24>=0.3,"Supported","Not supported (spread < 0.3 pts)")',
         "Centres perform alike. Fixes should be company-wide, not aimed at one site."),
        ("H5", "Chatbot gives lower CSAT than other channels",
         '="Chatbot "&TEXT(C18,"0.00")&" vs rest "&TEXT(C19,"0.00")', f"=B{h5_t}", f"=B{h5_p}",
         '=IF(E33<$B$3,IF(D33<0,"Supported","Rejected (opposite direction)"),"Not supported")',
         "Chatbot is slightly lower. Look at the chatbot hand-off and how it resolves issues."),
        ("H6", "Unhappy customers skip the survey",
         '="Response rate Very Neg "&TEXT(Pivots!E5,"0.0%")&" vs Very Pos "&TEXT(Pivots!E9,"0.0%")',
         "=Pivots!E5-Pivots!E9", "n/a",
         '=IF(ABS(D34)>=0.03,"Supported","Not supported (gap < 3 pts)")',
         "About 63% of all customers give no feedback, happy or not. The survey itself is a blind spot."),
    ]
    for i, v in enumerate(verdicts, start=29):
        for j, val in enumerate(v, start=1):
            c = ws.cell(i, j, val)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 4).number_format = "0.00"
        ws.cell(i, 5).number_format = "0.0000"
        ws.cell(i, 6).fill, ws.cell(i, 6).font = KPI_FILL, BOLD
    ws.cell(34, 4).number_format = "0.0%"
    box(ws, 29, 1, 34, 7)
    ws["A36"] = ("Statistic column: t statistic for H1, H2, H3, H5; CSAT spread in points for H4; response-rate gap for H6. "
                 "With n = 11,214 rated calls, even tiny differences can be 'significant'. Read the effect size as well as the p-value.")
    ws["A36"].font = NOTE_FONT
    widths(ws, {"A": 34, "B": 42, "C": 44, "D": 14, "E": 12, "F": 26, "G": 60})

    # ------------------------------------------------------------ Dashboard
    ws = wb.create_sheet("Dashboard")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Flipkart Customer Service Dashboard (Oct 2020)"
    ws["A1"].font = Font(name=FONT, bold=True, size=18, color="FFFFFF")
    for c in range(1, 19):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=BLUE)
    ws.row_dimensions[1].height = 30
    ws["A2"] = "Pick a filter in the yellow cells. Every KPI, table and chart on this sheet updates."
    ws["A2"].font = NOTE_FONT

    filters = [("Call centre", ["All"] + CENTERS), ("Channel", ["All"] + CHANNELS), ("Reason", ["All"] + REASONS)]
    for i, (lab, opts) in enumerate(filters, start=4):
        ws.cell(i, 1, lab).font = BOLD
        c = ws.cell(i, 2, "All")
        c.fill, c.border, c.font = INPUT_FILL, BOX, Font(name=FONT, bold=True, color="0000FF")
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=False)
        dv.add(c)
        ws.add_data_validation(dv)
        ws.cell(i, 3, f'=IF(B{i}="All","*",B{i})').font = Font(name=FONT, color="FFFFFF")  # hidden criteria
    F = f'{rng("center")},$C$4,{rng("channel")},$C$5,{rng("reason")},$C$6'

    kpi_defs = [
        ("Contacts", f"=COUNTIFS({F})", "#,##0"),
        ("Avg CSAT", f'=IFERROR(AVERAGEIFS({rng("csat")},{F}),"")', "0.00"),
        ("Response rate", f"=IFERROR(SUMIFS({rng('responded')},{F})/D9,0)", "0.0%"),
        ("% Satisfied", f'=IFERROR(COUNTIFS({F},{rng("csat")},">=8")/SUMIFS({rng("responded")},{F}),0)', "0.0%"),
        ("Neg. sentiment", f"=IFERROR(SUMIFS({rng('neg')},{F})/D9,0)", "0.0%"),
        ("SLA breach", f"=IFERROR(SUMIFS({rng('sla_breach')},{F})/D9,0)", "0.0%"),
        ("Avg duration", f"=IFERROR(AVERAGEIFS({rng('dur')},{F}),0)", "0.0"),
    ]
    for j, (lab, f, fmt) in enumerate(kpi_defs):
        col = 4 + j * 2
        lc_ = ws.cell(8, col, lab)
        lc_.font, lc_.fill = Font(name=FONT, bold=True, color="FFFFFF", size=9), PatternFill("solid", fgColor=NAVY)
        lc_.alignment = Alignment(horizontal="center")
        v = ws.cell(9, col, f)
        v.number_format, v.fill = fmt, KPI_FILL
        v.font, v.alignment = Font(name=FONT, bold=True, size=16, color=NAVY), Alignment(horizontal="center")
        ws.merge_cells(start_row=8, start_column=col, end_row=8, end_column=col + 1)
        ws.merge_cells(start_row=9, start_column=col, end_row=9, end_column=col + 1)
    ws.row_dimensions[9].height = 30
    ws.cell(10, 4, "vs overall CSAT:").font = NOTE_FONT
    ws.cell(10, 6, '=IFERROR(F9-Descriptive_Stats!B24,"")').number_format = '+0.00;-0.00;0.00'

    def dash_table(r, c, label, key, items):
        ws.cell(r - 1, c, f"By {label}").font = SUB_FONT
        for j, h in enumerate([label, "Calls", "Avg CSAT", "Resp. rate", "SLA breach"], start=c):
            ws.cell(r, j, h)
        style_header(ws, r, c, c + 4)
        L = get_column_letter(c)
        for i, it in enumerate(items, start=r + 1):
            ws.cell(i, c, it)
            k = rng(key)
            B = f"{get_column_letter(c+1)}{i}"
            ws.cell(i, c + 1, f"=COUNTIFS({F},{k},${L}{i})").number_format = "#,##0"
            ws.cell(i, c + 2, f'=IFERROR(AVERAGEIFS({rng("csat")},{F},{k},${L}{i}),0)').number_format = "0.00"
            ws.cell(i, c + 3, f"=IFERROR(SUMIFS({rng('responded')},{F},{k},${L}{i})/{B},0)").number_format = "0.0%"
            ws.cell(i, c + 4, f"=IFERROR(SUMIFS({rng('sla_breach')},{F},{k},${L}{i})/{B},0)").number_format = "0.0%"
        box(ws, r + 1, c, r + len(items), c + 4)
        return r + 1, r + len(items)

    s1 = dash_table(13, 1, "sentiment", "sentiment", SENTIMENTS)       # rows 14-18
    s2 = dash_table(21, 1, "response_time", "resp", RESPONSE)          # rows 22-24
    s3 = dash_table(27, 1, "duration_band", "dur_band", DUR_BANDS)     # rows 28-31
    s4 = dash_table(34, 1, "call_center", "center", CENTERS)           # rows 35-38

    def dchart(title_, rows, anchor, col=3, ytitle="Avg CSAT", pct=False):
        ch = BarChart()
        ch.title, ch.y_axis.title = title_, ytitle
        ch.add_data(Reference(ws, min_col=col, min_row=rows[0], max_row=rows[1]))
        ch.set_categories(Reference(ws, min_col=1, min_row=rows[0], max_row=rows[1]))
        ch.legend = None
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        ch.y_axis.delete = ch.x_axis.delete = False
        if pct:
            ch.y_axis.numFmt = "0%"
        ch.height, ch.width = 6.5, 12
        ws.add_chart(ch, anchor)

    dchart("Avg CSAT by sentiment", s1, "G12")
    dchart("Avg CSAT by response time", s2, "N12")
    dchart("Avg CSAT by call duration", s3, "G26")
    dchart("Avg CSAT by call centre", s4, "N26")
    dchart("Survey response rate by sentiment", s1, "G40", col=4, ytitle="Response rate", pct=True)
    dchart("SLA breach % by call centre", s4, "N40", col=5, ytitle="Above SLA %", pct=True)

    widths(ws, {"A": 16, "B": 14, "C": 10, "D": 11, "E": 11})
    for i in range(6, 19):
        ws.column_dimensions[get_column_letter(i)].width = 10

    # order sheets: README, Dashboard first for presentation
    order = ["README", "Dashboard", "Metrics_Hypotheses", "Hypothesis_Tests", "Descriptive_Stats", "Correlation",
             "Pivots", "Charts", "Cleaning_Log", "Clean_Data", "Raw_Data"]
    wb._sheets = [wb[s] for s in order]
    wb.active = 1
    wb.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
